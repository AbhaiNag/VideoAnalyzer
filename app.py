import os
import json
import base64
import tempfile
import shutil
from pathlib import Path
from typing import List, Optional

from fastapi import FastAPI, File, UploadFile, Form, HTTPException
from fastapi.staticfiles import StaticFiles
from fastapi.responses import HTMLResponse
import anthropic

try:
    import av
    HAS_AV = True
except ImportError:
    HAS_AV = False

try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

try:
    from PIL import Image
    import io
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

app = FastAPI(title="Video Analyzer Agent")
app.mount("/static", StaticFiles(directory="static"), name="static")

client = anthropic.Anthropic(api_key=os.environ.get("ANTHROPIC_API_KEY"))

SYSTEM_PROMPT = """You are an expert Technology Innovation Analyst and Venture Advisor specializing in identifying cutting-edge technology advancements with strong business viability.

Your mission is to analyze video content (provided as frame captures) and supporting presentation materials to evaluate whether the technology demonstrated represents a genuine cutting-edge advancement with meaningful business impact.

## Evaluation Framework

You assess content across four dimensions (each scored 0-25, total 100):

### 1. Technology Currency (0-25)
How current and advanced is the technology stack?
- 23-25: Uses frontier technologies (LLMs, quantum computing, advanced robotics, neuromorphic chips, etc.)
- 18-22: Uses modern but established technologies (cloud-native, ML/AI, edge computing, etc.)
- 12-17: Uses contemporary technologies with some innovation
- 6-11: Uses older technologies with minor improvements
- 0-5: Legacy or commodity technology with no advancement

### 2. Innovation & Novelty (0-25)
How novel and differentiated is the approach?
- 23-25: Breakthrough innovation, solves previously unsolved problems
- 18-22: Significant novelty, unique combination of existing technologies
- 12-17: Moderate innovation, improves on existing solutions
- 6-11: Incremental improvement, limited differentiation
- 0-5: No meaningful innovation, replication of existing solutions

### 3. Business Viability (0-25)
How strong is the business case and market opportunity?
- 23-25: Massive TAM, clear monetization, proven demand, strong ROI
- 18-22: Large market, viable business model, good ROI potential
- 12-17: Moderate market, reasonable business case
- 6-11: Limited market or unclear monetization
- 0-5: No clear business case or market

### 4. Technical Depth (0-25)
How technically rigorous and deep is the implementation?
- 23-25: Deep technical expertise demonstrated, novel algorithms or architectures
- 18-22: Strong technical foundation, good implementation depth
- 12-17: Adequate technical depth, some gaps
- 6-11: Shallow technical treatment, limited implementation detail
- 0-5: Superficial technical content

## Cutting-Edge Technology Landscape (2024-2025)
Correlate against:
- **AI/ML**: Large Language Models, Multimodal AI, Agentic AI, RAG, Fine-tuning, Edge AI, Neuromorphic computing
- **Cloud & Infrastructure**: Serverless, FinOps, Multi-cloud, Service Mesh, eBPF, WebAssembly
- **Data**: Real-time streaming, Vector databases, Data mesh, Lakehouse architecture
- **Security**: Zero Trust, Confidential computing, Post-quantum cryptography, AI-driven SOC
- **Hardware**: Custom silicon (TPUs/NPUs), Quantum processors, Advanced sensors, AR/VR hardware
- **Software Engineering**: Platform Engineering, GitOps, AI-assisted development, SRE
- **Emerging**: Spatial computing, Digital twins, Synthetic data, Federated learning

Provide detailed, evidence-based analysis grounded in what is actually visible in the content provided."""

ANALYSIS_SCHEMA = {
    "type": "object",
    "properties": {
        "executive_summary": {"type": "string"},
        "overall_score": {"type": "integer"},
        "recommendation": {
            "type": "string",
            "enum": [
                "Highly Recommended for Innovation",
                "Recommended for Innovation",
                "Conditionally Recommended",
                "Not Recommended"
            ]
        },
        "scores": {
            "type": "object",
            "properties": {
                "technology_currency": {"type": "integer"},
                "innovation_novelty": {"type": "integer"},
                "business_viability": {"type": "integer"},
                "technical_depth": {"type": "integer"}
            },
            "required": ["technology_currency", "innovation_novelty", "business_viability", "technical_depth"],
            "additionalProperties": False
        },
        "technology_stack": {"type": "array", "items": {"type": "string"}},
        "cutting_edge_correlation": {"type": "string"},
        "technology_analysis": {"type": "string"},
        "business_case": {"type": "string"},
        "key_innovations": {"type": "array", "items": {"type": "string"}},
        "business_opportunities": {"type": "array", "items": {"type": "string"}},
        "improvement_areas": {"type": "array", "items": {"type": "string"}},
        "market_positioning": {"type": "string"},
        "detailed_analysis": {"type": "string"}
    },
    "required": [
        "executive_summary", "overall_score", "recommendation", "scores",
        "technology_stack", "cutting_edge_correlation", "technology_analysis",
        "business_case", "key_innovations", "business_opportunities",
        "improvement_areas", "market_positioning", "detailed_analysis"
    ],
    "additionalProperties": False
}


def extract_video_frames(video_path: str, num_frames: int = 12) -> List[str]:
    if not HAS_AV:
        return []
    frames_b64 = []
    try:
        container = av.open(video_path)
        stream = container.streams.video[0]
        duration_seconds = float(stream.duration * stream.time_base) if stream.duration else None
        if duration_seconds is None:
            duration_seconds = float(container.duration / 1_000_000) if container.duration else 0
        if duration_seconds > 480:
            raise HTTPException(status_code=400, detail="Video exceeds 8-minute limit")
        if duration_seconds <= 0:
            duration_seconds = 60
        interval = duration_seconds / (num_frames + 1)
        seek_times = [interval * (i + 1) for i in range(num_frames)]
        for seek_time in seek_times:
            try:
                seek_pts = int(seek_time / float(stream.time_base))
                container.seek(seek_pts, stream=stream)
                for frame in container.decode(video=0):
                    img = frame.to_image()
                    if HAS_PIL:
                        img.thumbnail((960, 540), Image.LANCZOS)
                    buf = io.BytesIO()
                    img.save(buf, format="JPEG", quality=75)
                    frames_b64.append(base64.b64encode(buf.getvalue()).decode("utf-8"))
                    break
            except Exception:
                continue
        container.close()
    except HTTPException:
        raise
    except Exception as e:
        print(f"Frame extraction warning: {e}")
    return frames_b64


def extract_pptx_text(file_path: str) -> str:
    if not HAS_PPTX:
        return ""
    try:
        prs = Presentation(file_path)
        slides_text = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            if texts:
                slides_text.append(f"--- Slide {i} ---\n" + "\n".join(texts))
        return "\n\n".join(slides_text)
    except Exception as e:
        print(f"PPTX extraction warning: {e}")
        return ""


@app.get("/", response_class=HTMLResponse)
async def serve_index():
    index_path = Path("static/index.html")
    if index_path.exists():
        return HTMLResponse(content=index_path.read_text())
    return HTMLResponse(content="<h1>Video Analyzer</h1>")


@app.post("/api/analyze")
async def analyze(
    video: UploadFile = File(...),
    presentation: Optional[UploadFile] = File(None),
    industry: str = Form(default="General Technology"),
    perspective: str = Form(default="All"),
    focus_areas: str = Form(default=""),
    custom_question: str = Form(default=""),
    analysis_depth: str = Form(default="Standard")
):
    tmp_dir = tempfile.mkdtemp()
    try:
        video_path = os.path.join(tmp_dir, video.filename or "video.mp4")
        with open(video_path, "wb") as f:
            shutil.copyfileobj(video.file, f)

        pres_path = None
        if presentation and presentation.filename:
            pres_path = os.path.join(tmp_dir, presentation.filename)
            with open(pres_path, "wb") as f:
                shutil.copyfileobj(presentation.file, f)

        content = []

        # Analysis parameters context
        params_text = f"""## Analysis Parameters
- **Industry / Domain**: {industry}
- **Evaluation Perspective**: {perspective}
- **Focus Areas**: {focus_areas if focus_areas else 'All dimensions equally weighted'}
- **Analysis Depth**: {analysis_depth}
{f'- **Specific Question**: {custom_question}' if custom_question.strip() else ''}

Please tailor your analysis to these parameters — prioritize the specified industry context, apply the requested perspective, and if a specific question is provided, address it directly in the detailed_analysis field.
"""
        content.append({"type": "text", "text": params_text})

        # Video frames
        frames = extract_video_frames(video_path)
        if frames:
            content.append({
                "type": "text",
                "text": f"## Video Content\nExtracted {len(frames)} frames from the video at evenly-spaced intervals:"
            })
            for frame_b64 in frames:
                content.append({
                    "type": "image",
                    "source": {"type": "base64", "media_type": "image/jpeg", "data": frame_b64}
                })
        else:
            content.append({
                "type": "text",
                "text": f"## Video Content\nFile: {video.filename} (frame extraction unavailable)"
            })

        # Presentation
        if pres_path:
            ext = Path(pres_path).suffix.lower()
            if ext in (".pptx", ".ppt"):
                pptx_text = extract_pptx_text(pres_path)
                if pptx_text:
                    content.append({
                        "type": "text",
                        "text": f"## Supporting Presentation\n\n{pptx_text}"
                    })
            elif ext == ".pdf":
                with open(pres_path, "rb") as f:
                    pdf_data = base64.b64encode(f.read()).decode("utf-8")
                content.append({
                    "type": "document",
                    "source": {"type": "base64", "media_type": "application/pdf", "data": pdf_data}
                })

        content.append({
            "type": "text",
            "text": """## Task
Perform a comprehensive technology innovation analysis of the content above using the parameters specified. Score each dimension honestly based on evidence only. Calculate overall_score as the exact sum of the four dimension scores."""
        })

        response = client.messages.create(
            model="claude-opus-4-7",
            max_tokens=4096,
            system=SYSTEM_PROMPT,
            messages=[{"role": "user", "content": content}],
            output_config={
                "format": {
                    "type": "json_schema",
                    "json_schema": {
                        "name": "technology_analysis",
                        "schema": ANALYSIS_SCHEMA
                    }
                }
            }
        )

        result_text = next((b.text for b in response.content if b.type == "text"), "{}")
        result = json.loads(result_text)

        if "scores" in result:
            s = result["scores"]
            result["overall_score"] = (
                s.get("technology_currency", 0) + s.get("innovation_novelty", 0) +
                s.get("business_viability", 0) + s.get("technical_depth", 0)
            )

        return result

    except anthropic.BadRequestError as e:
        raise HTTPException(status_code=400, detail=f"Content rejected by Claude: {str(e)}")
    except anthropic.AuthenticationError:
        raise HTTPException(status_code=401, detail="Invalid Anthropic API key.")
    except anthropic.RateLimitError:
        raise HTTPException(status_code=429, detail="Rate limit reached. Please try again shortly.")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Analysis failed: {str(e)}")
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)
